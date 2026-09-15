#!/usr/bin/env python3
"""Esporta i listini a griglia (prezzi di vendita, senza vetro) in xlsx: un file per serie, un foglio per tipologia.
Stesse formule del programma Listini_Serie.html (coef/kitRighe/costo) con i parametri di consegna di dati_listini_serie.json."""
import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
HERE = os.path.dirname(os.path.abspath(__file__)); D = json.load(open(os.path.join(HERE, 'dati_listini_serie.json'), encoding='utf-8'))
r2 = lambda x: round(x * 100 + 1e-9) / 100
ev = lambda co, L, H: co[0] + co[1] * L + co[2] * H
def coef(S, t):
    tt = [0, 0, 0]; nn = [0, 0, 0]; gg = [0, 0, 0]; acc = 0
    for p in t['profili']:
        if p['kg_m'] is None: continue
        a = tt if p['classe'] == 'tt' else nn
        for i in range(3): a[i] += p['pz'] * (p['kg_m'] / 1000) * p['lin'][i]
    for g in t['guarn']:
        if g['pr'] is None: continue
        for i in range(3): gg[i] += (g['pr'] / 1000) * g['lin'][i]
    for a in t['acc']:
        if a['pr'] is not None: acc += a['pz'] * a['pr']
    return tt, nn, gg, acc
def kit_righe(S, t, L, H):
    kit = t['kit']; out = []; p = S['param']
    if kit['tipo'] != 'blk': return out
    lin = lambda e: __import__('re').fullmatch(r'[0-9LH+\-*/().]+', str(e).replace(' ', '')) and e
    def lin_ev(expr):
        s = str(expr).replace(',', '.').replace(' ', ''); import re
        s = re.sub(r'L/2', '(L/2)', s); s = re.sub(r'(\d)([LH(])', r'\1*\2', s); s = re.sub(r'\)([LH(\d])', r')*\1', s)
        return eval(s, {'__builtins__': {}}, {'L': L, 'H': H})
    aw = lin_ev(kit.get('anta_l') or 'L'); ah = lin_ev(kit.get('anta_h') or 'H')
    n_cern = next(q for lim, q in D.get('cerniere_per_anta', [[1300, 2], [2400, 3], [None, 4]]) if lim is None or ah <= lim)
    soglia = False
    for r in kit['righe']:
        if r.get('fascia_h') and not (r['fascia_h'][0] <= ah <= r['fascia_h'][1]): continue
        if r.get('fascia'):
            if r.get('fascia_h') is None and soglia: continue
            if not (r['fascia'][0] <= aw <= r['fascia'][1]): continue
            if r.get('fascia_h') is None: soglia = True
        pr = r['pr'] if r['pr'] is not None else 0
        sc = p['sc_acc'] if (r['pr'] is not None and 'listino AluK' in (r.get('fonte') or '')) else 0
        q = r['cern'] * n_cern if r.get('cern') else r['q']
        out.append((r['cod'], q, r2(pr * (1 - sc))))
    return out
def costo(S, t, c, L, H):
    p = S['param']; tt, nn, gg, acc = c
    kgT = ev(tt, L, H); kgN = ev(nn, L, H); gua = ev(gg, L, H)
    ferr = r2(sum(q * n for _, q, n in kit_righe(S, t, L, H)))
    co = r2(((kgT * (p['eur_kg_tt'] + p['add_kg']) + kgN * (p['eur_kg_n'] + p['add_kg'])) * (1 - p['sc_prof']) + (gua + acc) * (1 - p['sc_acc'])) * (1 + p['sfrido']) + ferr + t['ore'] * p['eur_h'])
    return co, r2(co * (1 + p['ricarico']))
def rng(a, b): return list(range(a, b + 1, 100))
def esporta(serie, out_dir):
    S = D['serie'][serie]; wb = Workbook(); wb.remove(wb.active); p = S['param']; oggi = datetime.date.today().strftime('%d/%m/%Y')
    idx = wb.create_sheet('Indice'); idx.append([f"LISTINI A GRIGLIA — {S['nome']} — SENZA VETRO — prezzi di vendita in €, IVA esclusa — {oggi}"]); idx.append([])
    idx.append(['Parametri: €/kg TT', p['eur_kg_tt'], '€/kg non isolati', p['eur_kg_n'], 'verniciatura €/kg', p['add_kg'], 'sconto profili', p['sc_prof'], 'sconto accessori', p['sc_acc'], 'sfrido', p['sfrido'], '€/h', p['eur_h'], 'ricarico', p['ricarico']])
    idx.append([]); idx.append(['Foglio', 'Tipologia', 'Variante', 'L min', 'L max', 'H min', 'H max', 'Ore']); idx['A1'].font = Font(bold=True, size=12)
    for k in S['ordine']:
        t = S['tip'][k]; c = coef(S, t); Ls = rng(*t['L']); Hs = rng(*t['H'])
        ws = wb.create_sheet(k[:31]); ws.append([f"LISTINO — {t['nome']} — {t.get('variante') or ''} — {S['nome']} — SENZA VETRO"]); ws['A1'].font = Font(bold=True)
        ws.append(['H \\ L'] + Ls)
        for cell in ws[2]: cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='E8E8E8'); cell.alignment = Alignment(horizontal='center')
        for H in Hs:
            ws.append([H] + [costo(S, t, c, L, H)[1] for L in Ls]); ws.cell(ws.max_row, 1).font = Font(bold=True)
            for j in range(2, len(Ls) + 2): ws.cell(ws.max_row, j).number_format = '#,##0.00'
        ws.append([]); ws.append([f"Prezzi di listino in €, IVA esclusa, senza vetro. Ore manodopera {t['ore']} h. Generato il {oggi}."])
        ws.column_dimensions['A'].width = 9
        for j in range(len(Ls)): ws.column_dimensions[ws.cell(2, j + 2).column_letter].width = 10
        ws.freeze_panes = 'B3'
        idx.append([k, t['nome'], t.get('variante') or '', t['L'][0], t['L'][1], t['H'][0], t['H'][1], t['ore']])
    for col, w in zip('ABCDEFGH', (10, 60, 60, 8, 8, 8, 8, 6)): idx.column_dimensions[col].width = w
    out = os.path.join(out_dir, f'Listini_{serie}.xlsx'); wb.save(out); return out, len(S['ordine'])
IWG = {'D67': ('LINEA IWG 67ID', 'IWG67ID', 'LINEA_IWG_67ID', 'Portoncini IWG', 60), 'D77': ('LINEA IWG 77ID', 'IWG77ID', 'LINEA_IWG_77ID', 'Portoncini IWG', 80), 'S140': ('IWG S140', 'IWG S140', 'LINEA_IWG_S140', 'Alzante Scorrevole IWG', 100)}
def prodotti_iwg(serie):
    """prodotti nel formato del preventivatore (catalogo_prodotti): stesso contenuto del pulsante nel programma HTML"""
    import re
    S = D['serie'][serie]; gamma, label, idp, fam, base = IWG[serie]; out = []
    for i, k in enumerate(S['ordine']):
        t = S['tip'][k]; c = coef(S, t); Ls = rng(*t['L']); Hs = rng(*t['H'])
        grid = {f'{H}x{L}': costo(S, t, c, L, H)[1] for H in Hs for L in Ls}
        fam_t = 'Scorrevoli' if serie == 'S140' and (t.get('gruppo') or '').startswith('S140R') else fam
        var_pulita = re.sub(r'\s*\[distinta[^\]]*\]', '', t.get('variante') or '')
        nome = f"{t['gruppo']} — {var_pulita}" if t.get('gruppo') and t.get('variante') else t['nome']
        out.append({'id': f'{idp}_{k}', 'materiale': 'Alluminio', 'famiglia': fam_t, 'gamma': gamma, 'gamma_label': label, 'code': k, 'name': nome, 'unit': 'cad.', 'dim1_label': 'Altezza', 'dim2_label': 'Larghezza', 'dim1_values': Hs, 'dim2_values': Ls, 'grid': grid, 'ordine': base + i})
    return out
def prodotti_iwg_c75s():
    import sys, io, contextlib
    sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'listini-c75s'))
    with contextlib.redirect_stdout(io.StringIO()):
        import genera_listini as g
    P = g.PARAM; NOMI = {'F1': 'Finestra 1 anta AR (C75S)', 'F2': 'Finestra 2 ante AR (C75S)', 'PF1': 'Portafinestra 1 anta AR (C75S)', 'PF2': 'Portafinestra 2 ante AR (C75S)', 'FISSO': 'Fisso (C75S)'}
    def listino(key, L, H):
        t = g.TIP[key]; (tC, tL, tH), (nC, nL, nH) = g.kg_coef(t['profili']); gC, gL, gH = g.g_coef(t['guarn']); acc = sum(g.PRZ[a] * q for a, q in t['acc'])
        co = round((((tC + tL * L + tH * H) * P['eur_kg_tt'] + (nC + nL * L + nH * H) * P['eur_kg_n']) * (1 - P['sc_prof']) + ((gC + gL * L + gH * H) + acc) * (1 - P['sc_acc'])) * (1 + P['sfrido']) + g.costo_ferr(key, L, H) + t['ore'] * P['eur_h'], 2)
        return round(co * (1 + P['ricarico']), 2)
    out = []
    for i, key in enumerate(('FISSO', 'F1', 'F2', 'PF1', 'PF2')):
        t = g.TIP[key]; Ls = rng(*t['L']); Hs = rng(*t['H'])
        out.append({'id': f'LINEA_IWG_75_C75S_{key}', 'materiale': 'Alluminio', 'famiglia': 'Battenti IWG', 'gamma': 'LINEA IWG 75', 'gamma_label': 'IWG75', 'code': f'C75S {key}', 'name': NOMI[key], 'unit': 'cad.', 'dim1_label': 'Altezza', 'dim2_label': 'Larghezza', 'dim1_values': Hs, 'dim2_values': Ls, 'grid': {f'{H}x{L}': listino(key, L, H) for H in Hs for L in Ls}, 'ordine': 50 + i})
    return out
def esporta_iwg(prod, out):
    """xlsx per impostazioni.html → Listini → "Import serie (xlsx)": foglio Prodotti (template serramenti) + un foglio griglia per codice"""
    wb = Workbook(); wb.remove(wb.active); ws = wb.create_sheet('Prodotti')
    ws.append(['IWG Template Serramenti']); ws.append([]); ws.append(['id', 'materiale', 'famiglia', 'gamma', 'gamma_label', 'code', 'name', 'unit', 'dim1_label', 'dim2_label', 'dim1_values', 'dim2_values', 'attivo', 'ordine'])
    for p in prod: ws.append([p['id'], p['materiale'], p['famiglia'], p['gamma'], p['gamma_label'], p['code'], p['name'], p['unit'], p['dim1_label'], p['dim2_label'], ','.join(map(str, p['dim1_values'])), ','.join(map(str, p['dim2_values'])), 'TRUE', p['ordine']])
    for col, w in zip('ABCDEFGHIJKLMN', (28, 12, 24, 18, 12, 12, 60, 7, 10, 10, 45, 45, 7, 7)): ws.column_dimensions[col].width = w
    for p in prod:
        g = wb.create_sheet(p['code'][:31]); g.append(['Altezza \\ Larghezza'] + p['dim2_values'])
        for H in p['dim1_values']: g.append([H] + [p['grid'][f'{H}x{L}'] for L in p['dim2_values']])
    wb.save(out); return out
if __name__ == '__main__':
    out_dir = os.path.join(HERE, 'dist'); os.makedirs(out_dir, exist_ok=True)
    for serie in D['serie']:
        f, n = esporta(serie, out_dir); print(f, n, 'tipologie')
        print(esporta_iwg(prodotti_iwg(serie), os.path.join(out_dir, f'IWG_{serie}.xlsx')), 'per il preventivatore')
    print(esporta_iwg(prodotti_iwg_c75s(), os.path.join(out_dir, 'IWG_C75S.xlsx')), 'per il preventivatore')
    # controllo
    S = D['serie']['D67']; t = S['tip']['P1IA67']; print('D67 P1IA67 1000x2200 ->', costo(S, t, coef(S, t), 1000, 2200))
    S = D['serie']['S140']; t = S['tip']['SXX140']; print('S140 SXX140 2900x2400 ->', costo(S, t, coef(S, t), 2900, 2400))
